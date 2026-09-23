"""Convert Excel surveys to the fixed Spark reader's unquoted || format."""
import argparse
import datetime
import os
import re
import tempfile

import openpyxl


DELIMITER = '||'
# Build a reusable detection pattern; this does not replace any text.
# The ranges cover control characters; the final entries cover Unicode line/
# paragraph separators and the ASCII double quote used by the CSV reader.
UNSAFE = re.compile(r'[\x00-\x1f\x7f-\x9f\u2028\u2029"]')


def sanitize_comment_value(value):
    if value is None:
        return ''
    text = str(value)
    # re.sub(pattern, replacement, text) replaces every matching piece of text.
    # r'...' keeps backslashes literal for the regular-expression engine.
    # [...] matches any listed character; + groups consecutive matches together.
    # \r = carriage return, \n = newline, \t = tab.
    # \x85 = Unicode next-line character; \u2028 = line separator;
    # \u2029 = paragraph separator. These can arrive in copied text.
    # Replace each consecutive group with ONE ordinary space, so a comment
    # stays on one output line. Example: 'Hello\r\nworld' becomes 'Hello world'.
    text = re.sub(r'[\r\n\t\x85\u2028\u2029]+', ' ', text)
    # Replace remaining control characters with spaces: \x00-\x1f includes
    # NUL, backspace and form feed; \x7f-\x9f includes DEL and C1 controls.
    # These are control codes, not ordinary letters, punctuation or emojis.
    # There is no + here: EACH matched character becomes one space.
    text = re.sub(r'[\x00-\x1f\x7f-\x9f]', ' ', text)
    # \| matches a literal pipe; {2,} means two or more consecutive pipes.
    # match.group() is the matched pipe sequence. ' '.join(...) inserts a
    # space between its characters: '|||' becomes '| | |'. No pipe is removed.
    # This prevents comment text from looking like the || column separator.
    # A single pipe inside a comment is left alone.
    text = re.sub(r'\|{2,}', lambda match: ' '.join(match.group()), text)
    # One input quote becomes exactly one output quote. Never CSV-unescape
    # doubled quotes or strip enclosing quotes: Excel already supplies cell text.
    # replace changes every ASCII double quote to the Unicode right double
    # quote (\u201d), so it is not interpreted as CSV quoting.
    # strip() then removes whitespace ONLY at the start/end, not quotes
    # or spaces between words. Existing curly quotes remain unchanged.
    text = text.replace('"', '\u201d').strip()
    # Add one space after a trailing pipe so it cannot join the next ||
    # separator. Example: field 'end|' becomes 'end| ' before joining fields.
    return text + ' ' if text.endswith('|') else text


def format_value(value, mode=None):
    if value is None:
        return ''
    if mode is None:
        return str(value)
    # Format codes: %Y = four-digit year, %m = month, %d = day,
    # %H = 24-hour hour, %M = minute, %S = second.
    target = '%Y-%m-%d %H:%M:%S' if mode == 'datetime' else '%Y-%m-%d'
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime(target)
    formats = (['%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M',
                '%m/%d/%Y %H:%M:%S', '%m/%d/%Y %H:%M']
               if mode == 'datetime' else ['%d/%m/%Y', '%m/%d/%Y'])
    for fmt in formats:
        try:
            # strptime reads the source date using fmt; strftime writes it
            # in the target BigQuery format. Day-first formats are tried first.
            return datetime.datetime.strptime(str(value), fmt).strftime(target)
        except ValueError:
            pass
    return str(value)


def resolve_target_column_indices(headers, names):
    result = set()
    for name in names:
        # strip removes surrounding whitespace; lower ignores capitalization
        # for matching only. This does not rename the Excel headers.
        matches = {i for i, value in enumerate(headers)
                   if value is not None and str(value).strip().lower() == name.strip().lower()}
        if not matches:
            print(f'Warning: column not found in header: {name}')
        result.update(matches)
    return result


def csv_from_excel(file_info, specific_text='No filters applied', date_columns=(), datetime_columns=()):
    """Sanitize comments; reject unsafe other fields without silently changing them.

    The fixed downstream reader uses ||, ASCII quote, and physical line records.
    CSV quoting cannot safely preserve embedded newlines under that contract.
    Publish the local output only after every row passes validation.
    """
    wb = openpyxl.load_workbook(file_info['input_file'], data_only=True)
    temporary_output = None
    try:
        sheet = wb.active
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        # Ignore surrounding whitespace/case when recognizing either comment
        # suffix. Only these columns receive the free-text replacements above.
        comments = {i for i, value in enumerate(headers)
                    if value is not None and str(value).strip().upper().endswith(('_COMMENT', '_COMMENTS'))}
        dates = resolve_target_column_indices(headers, date_columns)
        datetimes = resolve_target_column_indices(headers, datetime_columns)
        if dates & datetimes:
            print('Warning: datetime formatting takes precedence for overlapping date columns.')
        output_path = os.path.abspath(file_info['output_file'])
        count = 0
        pending_blank_rows = 0
        with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', newline='',
                                         dir=os.path.dirname(output_path), delete=False) as output:
            temporary_output = output.name
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                # Trim the first cell for the filter-marker comparison only;
                # the original value is still used when writing a normal row.
                first = '' if values[0] is None else str(values[0]).strip()
                if first.lower() == specific_text.lower():
                    continue
                fields = []
                for i, value in enumerate(values):
                    mode = 'datetime' if i in datetimes else 'date' if i in dates else None
                    text = format_value(value, mode)
                    if i in comments:
                        original_text = '' if value is None else str(value)
                        text = sanitize_comment_value(text)
                        # Compare with the actual Excel cell, not a CSV-rendered copy.
                        # Preserve existing curly quotes and every ASCII quote one-for-one.
                        # On a comparison copy, convert ASCII quotes one-for-one.
                        # findall collects every straight/left-curly/right-curly
                        # double quote in order; it does not edit the cell.
                        expected_quotes = re.findall(r'["\u201c\u201d]', original_text.replace('"', '\u201d'))
                        # Collect the output quotes too, to detect any loss,
                        # extra quotes, or change in the expected quote sequence.
                        actual_quotes = re.findall(r'["\u201c\u201d]', text)
                        if actual_quotes != expected_quotes:
                            cell = f'{sheet.title}!{openpyxl.utils.get_column_letter(i + 1)}{row_number}'
                            raise ValueError(f'Quote preservation failed at {cell}: expected '
                                             f'{len(expected_quotes)} quotes, got {len(actual_quotes)}. '
                                             'Original file retained.')
                    elif (UNSAFE.search(text) or DELIMITER in text
                          or (i < len(values) - 1 and text.endswith('|'))):
                        cell = f'{sheet.title}!{openpyxl.utils.get_column_letter(i + 1)}{row_number}'
                        raise ValueError(
                            f'Unsafe non-comment value at {cell}: line break, control character, '
                            'quote, or pipe would corrupt the output record. Correct the source '
                            'or explicitly enable cleanup for this column. Original file retained.'
                        )
                    fields.append(text)
                # Insert || between fields; splitting it again must recover
                # the exact same fields, or an embedded pipe caused ambiguity.
                record = DELIMITER.join(fields)
                if UNSAFE.search(record) or record.split(DELIMITER) != fields:
                    raise ValueError(f'Record validation failed at {sheet.title}, row {row_number}')
                # Preserve internal blank rows, but discard all trailing empty rows.
                if all(value == '' for value in fields):
                    pending_blank_rows += 1
                    continue
                if pending_blank_rows:
                    output.write((DELIMITER.join([''] * len(headers)) + '\n') * pending_blank_rows)
                    count += pending_blank_rows
                    pending_blank_rows = 0
                output.write(record + '\n')
                count += 1
        # This replaces the output FILE after validation; it is unrelated
        # to string replacement and does not change the file's contents.
        os.replace(temporary_output, output_path)
        temporary_output = None
        print(f'Processed file saved as: {output_path}; records: {count}')
        return count
    finally:
        wb.close()
        if temporary_output is not None:
            os.remove(temporary_output)


def main(args):
    # Local conversion and tests do not require GCS credentials or its SDK.
    from google.cloud import storage

    pattern = re.compile(args.file_name)
    # Split the argument at commas, trim each header name, and ignore empty
    # entries. Example: ' JOIN_DATE, EXIT_DATE, ' becomes two column names.
    date_columns = [c.strip() for c in args.date_columns.split(',') if c.strip()]
    # Apply the same comma-separated parsing to the datetime header list.
    datetime_columns = [c.strip() for c in args.datetime_columns.split(',') if c.strip()]
    landing_bucket = storage.Client(project=args.landing_project).bucket(args.landing_bucket)
    archive_bucket = storage.Client(project=args.staging_project).bucket(args.archive_bucket)
    jobs = []
    output_names = set()
    for blob in landing_bucket.list_blobs(prefix='SFG_Generic/'):
        if not pattern.search(blob.name) or not blob.name.lower().endswith(('.xlsx', '.xlsm')):
            continue
        # Take only the filename after the last / in the GCS object path.
        base = blob.name.rsplit('/', 1)[-1]
        # Replace the filename extension with .csv; this does not edit data.
        output_name = os.path.splitext(base)[0] + '.csv'
        if output_name in output_names:
            raise ValueError(f'Multiple input files map to output {output_name}; no files processed.')
        output_names.add(output_name)
        jobs.append((blob, output_name))
    if not jobs:
        raise ValueError(f'No Excel files matching {args.file_name!r} in SFG_Generic/.')

    for original, output_name in jobs:
        with tempfile.TemporaryDirectory(prefix='survey_preprocess_') as temp_dir:
            # Fixed local names avoid collisions and platform-specific object-name hazards.
            info = {'input_file': os.path.join(temp_dir, 'input.xlsx'),
                    'output_file': os.path.join(temp_dir, 'output.csv')}
            original.download_to_filename(info['input_file'], if_generation_match=original.generation)
            csv_from_excel(info, date_columns=date_columns, datetime_columns=datetime_columns)
            output_blob = landing_bucket.blob('SFG_Generic/' + output_name)
            output_blob.upload_from_filename(info['output_file'])
            # GCS copy must be invoked on the SOURCE bucket. Never delete the source
            # until both the validated output upload and archive copy have succeeded.
            destination = landing_bucket.copy_blob(
                original, archive_bucket,
                new_name='GLINT/ORGINAL/' + original.name.rsplit('/', 1)[-1],
                if_source_generation_match=original.generation)
            if destination is None:
                raise RuntimeError(f'Archive copy failed for {original.name}; original retained.')
            original.delete(if_generation_match=original.generation)
            print(f'Uploaded gs://{landing_bucket.name}/{output_blob.name} and archived {original.name}')
    print('All files processed and archived successfully.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Convert Excel surveys from GCS to validated ||-delimited records.')
    parser.add_argument('--landing_project', required=True, help='GCP project ID for landing bucket')
    parser.add_argument('--staging_project', required=True, help='GCP project ID for archive bucket')
    parser.add_argument('--landing_bucket', required=True)
    parser.add_argument('--archive_bucket', required=True)
    parser.add_argument('--file_name', required=True, help='Regular expression matching Excel object names')
    parser.add_argument('--date_columns', default='', help='Comma-separated headers to format as YYYY-MM-DD')
    parser.add_argument('--datetime_columns', default='', help='Comma-separated headers to format as YYYY-MM-DD HH:MM:SS')
    main(parser.parse_args())
